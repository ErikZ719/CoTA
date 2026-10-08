#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Render a one-slide .pptx to a vector PDF sized for the paper.

PowerPoint's own AppleScript export is unavailable on this machine, so the
slide is re-drawn from its XML with matplotlib.  The shape vocabulary used by
make_framework_pptx.py is covered: rectangles, rounded rectangles, ovals,
straight connectors (with arrow heads), freeform polylines, pictures and text
boxes, with solid or two-stop gradient fills.

Usage:  python pptx2pdf.py in.pptx out.pdf [width_cm]
"""
import sys
import textwrap
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, Ellipse, FancyBboxPatch
from matplotlib.transforms import Bbox
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

plt.rcParams['pdf.fonttype'] = 42
plt.rcParams['font.sans-serif'] = ['Helvetica Neue', 'Helvetica', 'Arial',
                                   'DejaVu Sans']
plt.rcParams['font.serif'] = ['Times New Roman', 'STIXGeneral', 'DejaVu Serif']
plt.rcParams['mathtext.fontset'] = 'stix'

src = sys.argv[1]
out = sys.argv[2]
width_cm = float(sys.argv[3]) if len(sys.argv) > 3 else 17.0

prs = Presentation(src)
SW = prs.slide_width / 914400.0
SH_ = prs.slide_height / 914400.0
W_in = width_cm / 2.54
H_in = W_in * SH_ / SW
fig = plt.figure(figsize=(W_in, H_in))
ax = fig.add_axes([0, 0, 1, 1])
ax.set_xlim(0, SW); ax.set_ylim(SH_, 0)
ax.set_aspect('equal'); ax.axis('off')
SCALE = W_in / SW                      # inches of output per inch of slide


def E(v):
    return v / 914400.0 if v is not None else 0.0


def hexc(el):
    c = el.find(qn('a:srgbClr'))
    return '#%s' % c.get('val') if c is not None else None


def fill_of(sh):
    spPr = sh._element.find(qn('p:spPr'))
    if spPr is None:
        return None
    gf = spPr.find(qn('a:gradFill'))
    if gf is not None:
        stops = [hexc(gs) for gs in gf.iter(qn('a:gs'))]
        stops = [s for s in stops if s]
        if stops:                       # average the two stops
            rgb = [int(sum(int(s[i:i + 2], 16) for s in stops) / len(stops))
                   for i in (1, 3, 5)]
            return '#%02x%02x%02x' % tuple(rgb)
    sf = spPr.find(qn('a:solidFill'))
    if sf is not None:
        return hexc(sf)
    return None


def line_of(sh):
    spPr = sh._element.find(qn('p:spPr'))
    if spPr is None:
        return None, 0.6
    ln = spPr.find(qn('a:ln'))
    if ln is None:
        return None, 0.6
    if ln.find(qn('a:noFill')) is not None:
        return None, 0.6
    sf = ln.find(qn('a:solidFill'))
    c = hexc(sf) if sf is not None else None
    w = ln.get('w')
    return c, (int(w) / 12700.0 if w else 0.75)


def geom(sh):
    spPr = sh._element.find(qn('p:spPr'))
    if spPr is None:
        return None, 0.0
    pg = spPr.find(qn('a:prstGeom'))
    if pg is None:
        return 'custom', 0.0
    adj = 0.0
    gd = pg.find(qn('a:avLst'))
    if gd is not None:
        g = gd.find(qn('a:gd'))
        if g is not None and g.get('fmla', '').startswith('val '):
            adj = int(g.get('fmla').split()[1]) / 100000.0
    return pg.get('prst'), adj


def draw(sh):
    x, y, w, h = E(sh.left), E(sh.top), E(sh.width), E(sh.height)
    st = str(sh.shape_type)

    if 'PICTURE' in st:
        img = sh.image.blob
        import io as _io
        from PIL import Image
        im = Image.open(_io.BytesIO(img))
        ax.imshow(im, extent=(x, x + w, y + h, y), zorder=3, aspect='auto')
        return

    if sh.__class__.__name__ == 'Connector' or 'LINE' in st:
        c, lw = line_of(sh)
        x1, y1 = E(sh.begin_x), E(sh.begin_y)
        x2, y2 = E(sh.end_x), E(sh.end_y)
        ln = sh._element.find(qn('p:spPr')).find(qn('a:ln'))
        dash = ln.find(qn('a:prstDash')) if ln is not None else None
        ls = (0, (2.4, 1.8)) if dash is not None else '-'
        head = ln.find(qn('a:tailEnd')) if ln is not None else None
        if head is not None:
            ax.annotate('', xy=(x2, y2), xytext=(x1, y1),
                        arrowprops=dict(arrowstyle='-|>', color=c or '#999',
                                        lw=lw * SCALE, shrinkA=0, shrinkB=0,
                                        linestyle=ls,
                                        mutation_scale=7 * SCALE))
        else:
            ax.plot([x1, x2], [y1, y2], color=c or '#999', lw=lw * SCALE, ls=ls,
                    solid_capstyle='round')
        return

    if 'FREEFORM' in st:
        c, lw = line_of(sh)
        spPr = sh._element.find(qn('p:spPr'))
        path = spPr.find('.//' + qn('a:path'))
        if path is None:
            return
        pw = float(path.get('w') or 1); ph = float(path.get('h') or 1)
        pts = []
        for node in path:
            pt = node.find(qn('a:pt'))
            if pt is not None:
                pts.append((float(pt.get('x')), float(pt.get('y'))))
        if len(pts) < 2:
            return
        xs = [x + (px / pw) * w for px, _ in pts]
        ys = [y + (py / ph) * h for _, py in pts]
        fc = fill_of(sh)
        if fc:
            ax.fill(xs, ys, color=fc, zorder=1, lw=0)
        if c:
            ax.plot(xs, ys, color=c, lw=lw * SCALE, solid_capstyle='round',
                    solid_joinstyle='round')
        return

    if sh.has_text_frame and (sh.text_frame.text or '').strip():
        runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
        s = sh.text_frame.text
        sz, c, bold, ital, fam = 9.0, '#333333', False, False, 'sans-serif'
        if runs:
            r0 = runs[0]
            if r0.font.size:
                sz = r0.font.size.pt
            try:
                if r0.font.color and r0.font.color.rgb:
                    c = '#%s' % str(r0.font.color.rgb)
            except Exception:
                pass
            bold = bool(r0.font.bold); ital = bool(r0.font.italic)
            fn = r0.font.name or ''
            fam = 'serif' if 'Times' in fn or 'STIX' in fn else 'sans-serif'
        if any(ch in s for ch in 'ℐ⌊⌋∖∪→↓⊤'):
            fam, ital = 'STIXGeneral', False
        al = str(sh.text_frame.paragraphs[0].alignment or '')
        ha, tx = 'left', x
        if 'CENTER' in al:
            ha, tx = 'center', x + w / 2
        elif 'RIGHT' in al:
            ha, tx = 'right', x + w
        rot = getattr(sh, 'rotation', 0) or 0
        body = s
        est = sz * 0.0073
        if len(s) * est > w and len(s) > 24:
            ncw = max(8, int(w / est))
            body = '\n'.join(textwrap.wrap(s, ncw))
        ax.text(tx, y + h / 2, body, ha=ha, va='center', fontsize=sz * SCALE,
                color=c, fontweight='bold' if bold else 'normal',
                style='italic' if ital else 'normal', rotation=rot, family=fam,
                zorder=6, linespacing=1.22)
        return

    fc = fill_of(sh)
    lc, lw = line_of(sh)
    prst, adj = geom(sh)
    if prst == 'ellipse':
        ax.add_patch(Ellipse((x + w / 2, y + h / 2), w, h, fc=fc or 'none',
                             ec=lc or 'none', lw=lw * SCALE, zorder=2))
    elif prst == 'roundRect':
        r = max(0.008, min(adj, 0.5) * min(w, h))
        ax.add_patch(FancyBboxPatch((x + r, y + r), max(w - 2 * r, 0.001),
                                    max(h - 2 * r, 0.001),
                                    boxstyle='round,pad=%g,rounding_size=%g'
                                             % (r, r),
                                    fc=fc or 'none', ec=lc or 'none',
                                    lw=lw * SCALE, zorder=2))
    else:
        ax.add_patch(Rectangle((x, y), w, h, fc=fc or 'none', ec=lc or 'none',
                               lw=lw * SCALE, zorder=2))


for sh in prs.slides[0].shapes:
    try:
        draw(sh)
    except Exception:
        pass

fig.savefig(out, dpi=600)
print('written:', out, '%.2f x %.2f cm' % (W_in * 2.54, H_in * 2.54))
