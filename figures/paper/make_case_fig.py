#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Case-study figure (Section VI-C), single column, drawn at its printed size.

    image + instruction
    LLaDA-V                      (no cache)
    + dLLM-Cache                 (the Repeat Curse)
    + dLLM-Cache + CoTA++        (ours)

IEEE column width is 3.5 in, so the canvas is 3.5 in wide and every font size below is the size on
paper: include the PDF with \\includegraphics[width=\\columnwidth]. Times New Roman, square hairline
frames, one accent colour (red) for the repeated tokens.

One layout, two outputs drawn by the same calls:
    images/fig-case-study.pdf (+ .png)      vector figure for the paper
    _tools/CoTA++_case_study_<tag>.pptx     editable PowerPoint slide of the same size

Run from the paper directory:   /opt/anaconda3/bin/python _tools/make_case_fig.py A
The two-row version with the per-component lenses is make_case_fig_full.py.
"""
import os, re, sys, importlib.util, warnings
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from matplotlib.textpath import TextPath
from matplotlib.font_manager import FontProperties
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

warnings.filterwarnings("ignore")
plt.rcParams["pdf.fonttype"] = 42
HERE = os.path.dirname(os.path.abspath(__file__)); PAPER = os.path.dirname(HERE)
CS = os.path.join(PAPER, "..", "..", "information_flow", "results", "case_study")
spec = importlib.util.spec_from_file_location("lenses", os.path.join(CS, "lenses.py"))
LZ = importlib.util.module_from_spec(spec); spec.loader.exec_module(LZ)
LZ.PREFIX = os.environ.get("CASE_PREFIX", "cs2")
TAG = sys.argv[1] if len(sys.argv) > 1 else "A"
IMAGES = {"H": "COCO_val2014_000000075990.jpg", "A": "COCO_val2014_000000118929.jpg",
          "B": "COCO_val2014_000000200231.jpg"}

FONT = "Times New Roman"
INK, RULE, NOTE, RED, RED_BG, HEAD_BG = "000000", "595959", "404040", "C00000", "FBE3E3", "F2F2F2"
GREEN = "1E7B34"                          # "no repetition"
FRAME_PT, RULE_PT = 0.9, 0.6              # outline and header rule, in points
MARGIN = 0.05                             # white space around the figure, in inches
W = 3.5                                   # inches = \columnwidth
BODY_PT, HEAD_PT, LH = 7.6, 8.0, 0.128   # body size, header size, line height (in)
PAD, N_LINES = 0.06, 5


def rgb(h): return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
def RGB(h): return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


_FP = {(b, i): FontProperties(family=FONT, weight="bold" if b else "normal", style="italic" if i else "normal")
       for b in (0, 1) for i in (0, 1)}
_TW = {}
def tw(s, size, bold=False, italic=False):
    k = (s, size, bold, italic)
    if k not in _TW:
        fp = _FP[(int(bold), int(italic))]
        if s.strip():
            _TW[k] = TextPath((0, 0), s, size=size, prop=fp).get_extents().width / 72.0
        else:
            _TW[k] = (TextPath((0, 0), "x x", size=size, prop=fp).get_extents().width
                      - TextPath((0, 0), "xx", size=size, prop=fp).get_extents().width) / 72.0
    return _TW[k]


# ------------------------------------------------------------------ text: mark, compact, wrap
def styled_words(text):
    """(word, style). A word is 'rep' when it repeats the word before it, or loops a piece of itself."""
    out, prev = [], ""
    for wd in text.replace("\n", " ").split():
        key = re.sub(r"\W", "", wd).lower()
        m = re.search(r"(.{2,8}?)\1{3,}", wd)
        if m:
            n = len(re.findall(re.escape(m.group(1)), wd)); head = wd[:m.start()]
            if head: out.append((head, None))
            out += [(m.group(1) * 4 + "…", "rep"), (f"(×{n})", "rep")]
        elif key and key == prev:
            if out and out[-1][1] != "rep": out[-1] = (out[-1][0], "rep")
            out.append((wd, "rep"))
        else:
            out.append((wd, None))
        prev = key
    return out


def compact(ws, keep=8):
    """a long run of one word keeps its first `keep` copies, then '… (×n)'"""
    out, i = [], 0
    while i < len(ws):
        j = i
        while j + 1 < len(ws) and ws[j][1] == "rep" and ws[j + 1][1] == "rep" and \
                re.sub(r"\W", "", ws[j + 1][0]).lower() == re.sub(r"\W", "", ws[i][0]).lower():
            j += 1
        n = j - i + 1
        out += (ws[i:i + keep] + [("…", "rep"), (f"(×{n})", "rep")]) if n > keep + 2 else ws[i:j + 1]
        i = j + 1
    return out


def wrap(ws, width, size, max_lines, jump=False):
    """greedy wrap. With jump, keep the longest opening that still leaves room for the whole loop."""
    if jump:
        first = next((k for k, (_, st) in enumerate(ws) if st == "rep"), None)
        if first is not None and len(wrap(ws, width, size, 10 ** 6)) > max_lines:
            tail = ws[max(first - 6, 0):]
            for h_ in range(max(first - 6, 0), 3, -1):
                cand = ws[:h_] + [("[…]", None)] + tail
                if len(wrap(cand, width, size, 10 ** 6)) <= max_lines:
                    ws = cand; break
    lines, cur, cw, sp = [], [], 0.0, tw(" ", size)
    for wd, st in ws:
        ww = tw(wd, size, st == "rep")
        if cur and cw + sp + ww > width:
            lines.append(cur); cur, cw = [], 0.0
            if len(lines) == max_lines: break
        cur.append((wd, st)); cw += (sp if len(cur) > 1 else 0) + ww
    else:
        if cur: lines.append(cur)
        return lines
    last = lines[-1]
    while last and sum(tw(a, size, s_ == "rep") for a, s_ in last) + sp * len(last) + tw("…", size) > width: last.pop()
    last.append(("…", None)); return lines


# ------------------------------------------------------------------ data and geometry
D = {a: LZ.load(TAG, a) for a in ("van", "cache", "full")}
longest = lambda a: max([0] + LZ.RM.run_lengths(LZ.RM.strip_layout(LZ.RM.trim_at_eot(list(D[a]["ids"]), 126081))))
ROWS = [("van", "LLaDA-V", True), ("cache", "+ dLLM-Cache", False), ("full", "+ dLLM-Cache + CoTA++", False)]
CW_ = W - 2 * MARGIN                       # content width
X0 = Y0 = MARGIN
TEXT_W = CW_ - 2 * PAD - 0.02
LINES = {a: wrap(compact(styled_words(D[a]["text"])), TEXT_W, BODY_PT, N_LINES, jump=(a == "cache")) for a, _, _ in ROWS}
IMG_W = 1.30
_im = Image.open(os.path.join(CS, IMAGES[TAG])).convert("RGB"); IMG_H = IMG_W * _im.size[1] / _im.size[0]
HEAD_H, GAP = 0.165, 0.055
BOX_H = {a: HEAD_H + len(LINES[a]) * LH + 0.05 for a, _, _ in ROWS}
H = 2 * MARGIN + IMG_H + GAP + sum(BOX_H.values()) + GAP * (len(ROWS) - 1)

fig = plt.figure(figsize=(W, H)); ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, W); ax.set_ylim(H, 0); ax.axis("off")
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(W), Inches(H)
SH = prs.slides.add_slide(prs.slide_layouts[6]).shapes
Z = [1]
def z():
    Z[0] += 1; return Z[0]


def frame(x, y, w, h, lc=RULE, lw=FRAME_PT, fc=None):
    ax.add_patch(Rectangle((x, y), w, h, fc=rgb(fc) if fc else "none", ec=rgb(lc) if lc else "none", lw=lw, zorder=z()))
    s = SH.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fc: s.fill.solid(); s.fill.fore_color.rgb = RGB(fc)
    else: s.fill.background()
    if lc: s.line.color.rgb = RGB(lc); s.line.width = Pt(lw)
    else: s.line.fill.background()
    s.shadow.inherit = False


def rule(x1, y1, x2, y2, color=RULE, lw=RULE_PT):
    ax.plot([x1, x2], [y1, y2], color=rgb(color), lw=lw, zorder=z(), solid_capstyle="butt")
    c = SH.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGB(color); c.line.width = Pt(lw)


def label(x, y, w, h, s, size, color=INK, bold=False, italic=False, align="left"):
    tx = {"left": x, "right": x + w, "center": x + w / 2}[align]
    ax.text(tx, y + h / 2, s, fontsize=size, color=rgb(color), family=FONT, ha=align, va="center",
            fontweight="bold" if bold else "normal", fontstyle="italic" if italic else "normal", zorder=z())
    tb = SH.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[align]
    r = p.add_run(); r.text = s; f = r.font; f.size = Pt(size); f.name = FONT; f.bold = bold; f.italic = italic
    f.color.rgb = RGB(color)


def words(x, y, w, lines):
    for i, ln in enumerate(lines):
        cx, cy = x, y + i * LH
        for wd, st in ln:
            ww = tw(wd, BODY_PT, st == "rep")
            if st == "rep":
                ax.add_patch(Rectangle((cx - 0.006, cy + 0.012), ww + 0.012, LH - 0.02, fc=rgb(RED_BG), ec="none", zorder=z()))
            ax.text(cx, cy + LH / 2, wd, fontsize=BODY_PT, family=FONT, va="center", ha="left", zorder=z(),
                    color=rgb(RED if st == "rep" else INK), fontweight="bold" if st == "rep" else "normal")
            cx += ww + tw(" ", BODY_PT)
    tb = SH.add_textbox(Inches(x), Inches(y), Inches(w), Inches(LH * len(lines))); tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.line_spacing = Pt(LH * 72)
        for j, (wd, st) in enumerate(ln):
            r = p.add_run(); r.text = wd + (" " if j < len(ln) - 1 else ""); f = r.font
            f.size = Pt(BODY_PT); f.name = FONT; f.bold = st == "rep"; f.color.rgb = RGB(RED if st == "rep" else INK)
            if st == "rep":
                rPr = r._r.get_or_add_rPr(); hl = rPr.makeelement(qn("a:highlight"), {})
                hl.append(rPr.makeelement(qn("a:srgbClr"), {"val": RED_BG}))
                rPr.insert_element_before(hl, "a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill", "a:latin", "a:ea", "a:cs",
                                          "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")


# ------------------------------------------------------------------ top strip: image and instruction
tmp = os.path.join(HERE, f"_case_{TAG}.jpg"); _im.save(tmp, quality=95)
ax.imshow(np.asarray(_im), extent=(X0, X0 + IMG_W, Y0 + IMG_H, Y0), zorder=z(), aspect="auto")
SH.add_picture(tmp, Inches(X0), Inches(Y0), Inches(IMG_W), Inches(IMG_H))
frame(X0, Y0, IMG_W, IMG_H)
tx0, tw0 = X0 + IMG_W + 0.10, CW_ - IMG_W - 0.10
label(tx0, Y0 + 0.015, tw0, 0.14, "Instruction", HEAD_PT, bold=True)
label(tx0, Y0 + 0.150, tw0, 0.14, "“Please describe the image in detail.”", BODY_PT, italic=True)
ly = Y0 + IMG_H - LH - 0.005
ax.add_patch(Rectangle((tx0, ly + 0.012), 0.205, LH - 0.02, fc=rgb(RED_BG), ec="none", zorder=z()))
_s = SH.add_shape(MSO_SHAPE.RECTANGLE, Inches(tx0), Inches(ly + 0.012), Inches(0.205), Inches(LH - 0.02))
_s.fill.solid(); _s.fill.fore_color.rgb = RGB(RED_BG); _s.line.fill.background(); _s.shadow.inherit = False
label(tx0 + 0.012, ly, 0.2, LH, "the", BODY_PT, color=RED, bold=True)
label(tx0 + 0.255, ly, tw0 - 0.255, LH, "token repeating the one before it", BODY_PT, color=NOTE)

# ------------------------------------------------------------------ the three responses
y = Y0 + IMG_H + GAP
for arm, name, uncached in ROWS:
    h = BOX_H[arm]
    frame(X0, y, CW_, HEAD_H, lc=None, fc=HEAD_BG)
    frame(X0, y, CW_, h); rule(X0, y + HEAD_H, X0 + CW_, y + HEAD_H)
    label(X0 + PAD, y, 2.2, HEAD_H, name, HEAD_PT, bold=True)
    n = longest(arm)
    note = ("longest repeated run: %d tokens" % n) if n >= 2 else "no repetition"
    label(X0 + CW_ - PAD - 1.9, y, 1.9, HEAD_H, note, BODY_PT, color=RED if n >= 2 else GREEN, italic=True,
          bold=(n < 2), align="right")
    words(X0 + PAD, y + HEAD_H + 0.025, TEXT_W, LINES[arm])
    y += h + GAP

fig.savefig(os.path.join(PAPER, "images", "fig-case-study.pdf"))
fig.savefig(os.path.join(PAPER, "images", "fig-case-study.png"), dpi=400)
prs.save(os.path.join(HERE, f"CoTA++_case_study_{TAG}.pptx"))
print("written: images/fig-case-study.pdf/.png (%.2f x %.2f in) and _tools/CoTA++_case_study_%s.pptx" % (W, H, TAG))
