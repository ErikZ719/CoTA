#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Case-study figure (Section VI-B), two columns, drawn at its printed size.

    (a) LLaDA-V with dLLM-Cache          (b) LaViDa with SlowFast
        image + instruction                  image + instruction
        LLaDA-V                              LaViDa
        + dLLM-Cache                         + SlowFast
        + dLLM-Cache + CoTA++                + SlowFast + CoTA++

Panel (a) is the single-column figure of make_case_fig.py, drawn by the same calls and unchanged. Panel (b) has the
same geometry. Its responses are the Table VI runs (lavida/t5/{van,sf,sfcpp}128_*), decoded token by token
(information_flow/results/case_study/lavida_sf128_cases.json), so that a marked token is one that belongs to a run
of identical content tokens, the criterion of ARR and MRL.

The canvas is the IEEE text width (two columns of 3.5 in and the 1 pc gutter): include the PDF with
\\includegraphics[width=\\textwidth]. Every font size below is the size on paper.

    images/fig-case-study-2col.pdf (+ .png)       vector figure for the paper
    _tools/CoTA++_case_study_2col_<case>.pptx     editable PowerPoint slide of the same size

Run from the paper directory:   /opt/anaconda3/bin/python _tools/make_case_fig_2col.py horse
                                /opt/anaconda3/bin/python _tools/make_case_fig_2col.py plate --preview
(--preview writes to _tools/preview_case_2col_<case>.pdf/.png instead of images/)
"""
import os, re, sys, json, importlib.util, warnings
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from matplotlib.textpath import TextPath
from matplotlib.font_manager import FontProperties
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

warnings.filterwarnings("ignore")
plt.rcParams["pdf.fonttype"] = 42
HERE = os.path.dirname(os.path.abspath(__file__)); PAPER = os.path.dirname(HERE)
CS = os.path.join(PAPER, "..", "..", "information_flow", "results", "case_study")
spec = importlib.util.spec_from_file_location("lenses", os.path.join(CS, "lenses.py"))
LZ = importlib.util.module_from_spec(spec); spec.loader.exec_module(LZ)
LZ.PREFIX = os.environ.get("CASE_PREFIX", "cs2")
ARGS = [a for a in sys.argv[1:] if not a.startswith("--")]
CASE = ARGS[0] if ARGS else "horse"
PREVIEW = "--preview" in sys.argv
IMAGE_A = "COCO_val2014_000000118929.jpg"
# crop = (left, top, right, bottom) in pixels, at the aspect ratio of the image of panel (a), 640 x 281
CASES = {"horse": dict(image="COCO_val2014_000000079380.jpg", crop=(0, 24, 640, 305), focus="arat"),
         "plate": dict(image="COCO_val2014_000000520862.jpg", crop=(0, 330, 480, 541), focus=None),
         # 2026-10-01 (user): panel (b) = the residual failure of the Discussion (image 536426, LLaDA-V, L=128), drawn from
         # fail_cases.json (keys hyd_van_128 / hyd_cache_128 / hyd_cpp_128); the LaViDa case moves to the appendix.
         "hydrant": dict(image="COCO_val2014_000000536426.jpg", crop=(0, 168, 375, 333), focus="hydrantrant", src="fail", prefix="hyd_",
                         names=(("van", "LLaDA-V"), ("cache", "+ dLLM-Cache"), ("full", "+ dLLM-Cache + CoTA++")),
                         sub_a="(a) LLaDA-V with dLLM-Cache: the typical case", sub="(b) LLaDA-V with dLLM-Cache: a residual failure")}
# 2026-10-01 (user): candidates for panel (b) from case_b_export.json (Table VI runs, decoded on the server):
#   bus   LLaDA-V, L=512, image 222340: the uncached model reads the number "33-41" on the bus (digits, blue), the cache
#         collapses into 341 commas, CoTA++ reads the number again;   mm512/mm128  MMaDA, image 534127, "pointed pointed".
CASES["bus"] = dict(image="COCO_val2014_000000222340.jpg", crop=(0, 150, 640, 431), focus="33", src="export", prefix="bus", L=512,
                    names=(("van", "LLaDA-V"), ("cache", "+ dLLM-Cache"), ("full", "+ dLLM-Cache + CoTA++")),
                    sub_a="(a) LLaDA-V with dLLM-Cache, $L{=}128$", sub="(b) LLaDA-V with dLLM-Cache, $L{=}512$")
for _L in (512, 128):
    CASES["mm%d" % _L] = dict(image="COCO_val2014_000000534127.jpg", crop=(0, 60, 640, 341), focus="pointed", src="export", prefix="mm", L=_L,
                              names=(("van", "MMaDA"), ("cache", "+ dLLM-Cache"), ("full", "+ dLLM-Cache + CoTA++")),
                              sub_a="(a) LLaDA-V with dLLM-Cache, $L{=}128$", sub="(b) MMaDA with dLLM-Cache, $L{=}%d$" % _L)
# residual-failure candidates with a broken cache and a harmless residual (a doubled period), LLaDA-V, L=512:
for _k, _im, _crop in (("cloth", "COCO_val2014_000000295491.jpg", (0, 60, 640, 341)), ("snow", "COCO_val2014_000000010142.jpg", (0, 100, 640, 381)),
                       ("road", "COCO_val2014_000000166509.jpg", (0, 80, 640, 361))):
    CASES[_k] = dict(image=_im, crop=_crop, focus="..", src="export", prefix=_k, L=512,
                     names=(("van", "LLaDA-V"), ("cache", "+ dLLM-Cache"), ("full", "+ dLLM-Cache + CoTA++")),
                     sub_a="(a) LLaDA-V with dLLM-Cache, L = 128", sub="(b) LLaDA-V with dLLM-Cache, L = 512")
ONLY_B = "--only-b" in sys.argv          # single-column figure of panel (b) alone, no sub-caption (appendix)

FONT = "Times New Roman"
INK, RULE, NOTE, RED, RED_BG, HEAD_BG = "000000", "595959", "404040", "C00000", "FBE3E3", "F2F2F2"
GREEN = "1E7B34"                          # "no repetition"
BLUE, BLUE_BG = "1F4E9C", "DCE6F5"           # digits of a number written in the image (not a repetition)
STYLE = {"rep": (RED, RED_BG), "num": (BLUE, BLUE_BG)}
FRAME_PT, RULE_PT = 0.9, 0.6              # outline and header rule, in points
MARGIN = 0.05                             # white space around a panel, in inches
WP = 3.5                                  # one panel = \columnwidth
SEP = 1.0 / 6.0                           # \columnsep of IEEEtran, 1 pc
W = WP if ONLY_B else 2 * WP + SEP        # \textwidth, or one column for --only-b
BODY_PT, HEAD_PT, LH = 7.6, 8.0, 0.128   # body size, header size, line height (in)
SUB_PT, SUB_H = 8.0, 0.17                 # "(a) ..." under each panel
PAD, N_LINES = 0.06, 5


def rgb(h): return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
def RGB(h): return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


_FP = {(b, i): FontProperties(family=FONT, weight="bold" if b else "normal", style="italic" if i else "normal")
       for b in (0, 1) for i in (0, 1)}
_TW = {}
def tw(s, size, bold=False, italic=False):
    k = (s, size, bold, italic)
    if k not in _TW:
        fp = _FP[(int(bold), int(italic))]
        if s.strip():
            _TW[k] = TextPath((0, 0), s, size=size, prop=fp).get_extents().width / 72.0
        else:
            _TW[k] = (TextPath((0, 0), "x x", size=size, prop=fp).get_extents().width
                      - TextPath((0, 0), "xx", size=size, prop=fp).get_extents().width) / 72.0
    return _TW[k]


# A word is a list of segments (text, style): the segments of one word are drawn without a space between them.
def ww_(word, size): return sum(tw(t, size, st == "rep") for t, st in word)
def is_rep(word): return any(st == "rep" for _, st in word)


# ------------------------------------------------------------------ panel (a): mark, compact (as make_case_fig.py)
def styled_words(text):
    """(word, style). A word is 'rep' when it repeats the word before it, or loops a piece of itself."""
    out, prev = [], ""
    for wd in text.replace("\n", " ").split():
        key = re.sub(r"\W", "", wd).lower()
        m = re.search(r"(.{2,8}?)\1{3,}", wd)
        if m:
            n = len(re.findall(re.escape(m.group(1)), wd)); head = wd[:m.start()]
            if head: out.append((head, None))
            out += [(m.group(1) * 4 + "…", "rep"), (f"(×{n})", "rep")]
        elif key and key == prev:
            if out and out[-1][1] != "rep": out[-1] = (out[-1][0], "rep")
            out.append((wd, "rep"))
        else:
            out.append((wd, None))
        prev = key
    return out


def compact(ws, keep=8):
    """a long run of one word keeps its first `keep` copies, then '… (×n)'"""
    out, i = [], 0
    while i < len(ws):
        j = i
        while j + 1 < len(ws) and ws[j][1] == "rep" and ws[j + 1][1] == "rep" and \
                re.sub(r"\W", "", ws[j + 1][0]).lower() == re.sub(r"\W", "", ws[i][0]).lower():
            j += 1
        n = j - i + 1
        out += (ws[i:i + keep] + [("…", "rep"), (f"(×{n})", "rep")]) if n > keep + 2 else ws[i:j + 1]
        i = j + 1
    return out


# ------------------------------------------------------------------ panel (b): words from the decoded tokens
def token_words(pieces, ids, keep_words=8, keep_pieces=4):
    """Words of a response, from its tokens. A token is 'rep' when it belongs to a run of identical content tokens
    (layout tokens are skipped, as in the metric). A run longer than keep + 2 is shortened to '… (×n)'."""
    end = next((k for k, p in enumerate(pieces) if p.startswith("<|")), len(pieces))
    pieces, ids = pieces[:end], ids[:end]
    content = [k for k, t in enumerate(ids) if t not in LZ.RM.LAYOUT_IDS]
    run_of, i = {}, 0
    while i < len(content):
        j = i
        while j + 1 < len(content) and ids[content[j + 1]] == ids[content[i]]: j += 1
        if j > i:
            for k in content[i:j + 1]: run_of[k] = (content[i], j - i + 1)
        i = j + 1
    segs, k = [], 0                       # (text, style), text may start with a space
    while k < len(pieces):
        if k in run_of and run_of[k][0] == k:
            n = run_of[k][1]; p = pieces[k]; last = [q for q in content if q >= k][n - 1]
            st = "num" if p.strip().isdigit() else "rep"
            if p.startswith(" "):         # a run of whole words
                shown = n if n <= keep_words + 2 else keep_words
                segs += [(p, st)] * shown
                if shown < n: segs += [(" …", st), (" (×%d)" % n, st)]
            else:                         # a run inside one word
                shown = n if n <= keep_pieces + 2 else keep_pieces
                segs.append((p * shown + ("…" if shown < n else ""), st))
                if shown < n: segs.append((" (×%d)" % n, st))
            k = last + 1
        else:
            segs.append((pieces[k].replace("\n", " "), None)); k += 1
    words, cur = [], []
    for t, st in segs:
        parts = re.split(r"(\s+)", t)
        for part in parts:
            if not part: continue
            if part.isspace():
                if cur: words.append(cur); cur = []
            elif cur and cur[-1][1] == st: cur[-1] = (cur[-1][0] + part, st)
            else: cur.append((part, st))
    if cur: words.append(cur)
    return words


# ------------------------------------------------------------------ wrap, with an elision that keeps the focus in view
ELLIPSIS = [("[…]", None)]
def wrap(ws, width, size, max_lines):
    lines, cur, cw, sp = [], [], 0.0, tw(" ", size)
    for word in ws:
        w_ = ww_(word, size)
        if cur and cw + sp + w_ > width:
            lines.append(cur); cur, cw = [], 0.0
            if len(lines) == max_lines: break
        cur.append(word); cw += (sp if len(cur) > 1 else 0) + w_
    else:
        if cur: lines.append(cur)
        return lines, False
    last = lines[-1]
    while last and sum(ww_(a, size) for a in last) + sp * len(last) + tw("…", size) > width: last.pop()
    last.append([("…", None)]); return lines, True


def n_lines(ws, width, size): return len(wrap(ws, width, size, 10 ** 6)[0])


def fit(ws, width, size, max_lines, focus=None, sentence=True):
    """ws wrapped to max_lines. If the words [focus[0], focus[1]) would not be shown in full, the longest opening
    that leaves room for them is kept, then '[…]'. The opening ends on a sentence when it can."""
    if focus is None or n_lines(ws[:focus[1]], width, size) <= max_lines:
        return wrap(ws, width, size, max_lines)[0]
    a, b = focus
    ends = [k + 1 for k in range(a) if ws[k][-1][0].endswith((".", "!", "?"))] if sentence else []
    for cands in (sorted(ends, reverse=True), range(a - 1, 3, -1)):
        for h_ in cands:
            if n_lines(ws[:h_] + [ELLIPSIS] + ws[a:b], width, size) <= max_lines:
                return wrap(ws[:h_] + [ELLIPSIS] + ws[a:], width, size, max_lines)[0]
    return wrap(ws, width, size, max_lines)[0]


def wrap_a(ws, width, size, max_lines, jump=False):
    """the wrap of make_case_fig.py, on (word, style) pairs: panel (a) must come out as in the single-column figure"""
    one = lambda W_: wrap([[x] for x in W_], width, size, 10 ** 6)[0]
    if jump:
        first = next((k for k, (_, st) in enumerate(ws) if st == "rep"), None)
        if first is not None and len(one(ws)) > max_lines:
            tail = ws[max(first - 6, 0):]
            for h_ in range(max(first - 6, 0), 3, -1):
                cand = ws[:h_] + [("[…]", None)] + tail
                if len(one(cand)) <= max_lines:
                    ws = cand; break
    return wrap([[x] for x in ws], width, size, max_lines)[0]


# ------------------------------------------------------------------ data
CW_ = WP - 2 * MARGIN                      # content width of a panel
TEXT_W = CW_ - 2 * PAD - 0.02
DA = {a: LZ.load("A", a) for a in ("van", "cache", "full")}
longest_a = lambda a: max([0] + LZ.RM.run_lengths(LZ.RM.strip_layout(LZ.RM.trim_at_eot(list(DA[a]["ids"]), 126081))))
PANEL_A = dict(
    sub="(a) LLaDA-V with dLLM-Cache", image=Image.open(os.path.join(CS, IMAGE_A)).convert("RGB"),
    rows=[(name, wrap_a(compact(styled_words(DA[a]["text"])), TEXT_W, BODY_PT, N_LINES, jump=(a == "cache")), longest_a(a))
          for a, name in (("van", "LLaDA-V"), ("cache", "+ dLLM-Cache"), ("full", "+ dLLM-Cache + CoTA++"))])

C = CASES[CASE]
if C.get("src") == "export":
    _D = json.load(open(os.path.join(CS, "case_b_export.json"))); _px, _L = C["prefix"], C["L"]
    DB = {"van": _D["%s_van_%d" % (_px, _L)], "cache": _D["%s_cache_%d" % (_px, _L)], "full": _D["%s_cpp_%d" % (_px, _L)]}
elif C.get("src") == "fail":
    _D = json.load(open(os.path.join(CS, "fail_cases.json"))); _px = C["prefix"]
    DB = {"van": _D[_px + "van_128"], "cache": _D[_px + "cache_128"], "full": _D[_px + "cpp_128"]}
else:
    DB = json.load(open(os.path.join(CS, "lavida_sf128_cases.json")))[C["image"]]
NAMES_B = C.get("names", (("van", "LaViDa"), ("cache", "+ SlowFast"), ("full", "+ SlowFast + CoTA++")))
def rows_b():
    out = []
    for a, name in NAMES_B:
        ws = token_words(DB[a]["pieces"], DB[a]["ids"]); focus = None
        reps = [k for k, w_ in enumerate(ws) if is_rep(w_)]
        if a == "cache" and reps:                                   # the longest run must be in view
            k0 = max(reps, key=lambda k: max(len(t) for t, st in ws[k] if st == "rep"))
            k0 = next(k for k in reps if any("×" in t for t, _ in ws[k])) if any("×" in t for k in reps for t, _ in ws[k]) else k0
            first = k0
            while first - 1 in reps: first -= 1
            focus = (max(first - 1, 0), min(k0 + 2, len(ws)))
        elif C["focus"]:                                            # the sentence that names the focus word
            hit = [k for k, w_ in enumerate(ws) if C["focus"] in "".join(t for t, _ in w_)]
            if hit:
                s0 = hit[0]
                while s0 > 0 and not ws[s0 - 1][-1][0].endswith((".", "!", "?")): s0 -= 1
                s1 = hit[0]
                while s1 < len(ws) - 1 and not ws[s1][-1][0].endswith((".", "!", "?")): s1 += 1
                focus = (s0, s1 + 1)
        best, digit = run_stats(DB[a]["ids"], DB[a]["pieces"])
        out.append((name, fit(ws, TEXT_W, BODY_PT, N_LINES, focus), best, digit))
    return out
def run_stats(ids, pieces):
    """longest run of non-digit content tokens, and whether a digit run exists (the digits of a number in the image)"""
    end = next((k for k, p in enumerate(pieces) if p.startswith("<|")), len(pieces)); ids, pieces = ids[:end], pieces[:end]
    best, cur, prev, digit = 0, 0, None, False
    for t, p in zip(ids, pieces):
        if t in LZ.RM.LAYOUT_IDS: continue
        cur = cur + 1 if t == prev else 1; prev = t
        if cur >= 2:
            if p.strip().isdigit(): digit = True
            else: best = max(best, cur)
    return best, digit
PANEL_A["sub"] = C.get("sub_a", PANEL_A["sub"])
PANEL_B = dict(sub="" if ONLY_B else C.get("sub", "(b) LaViDa with SlowFast"),
               image=Image.open(os.path.join(CS, C["image"])).convert("RGB").crop(C["crop"]), rows=rows_b())

IMG_W = 1.30
IMG_H = IMG_W * PANEL_A["image"].size[1] / PANEL_A["image"].size[0]
HEAD_H, GAP = 0.165, 0.055
def box_h(lines): return HEAD_H + len(lines) * LH + 0.05
for P in (PANEL_A, PANEL_B):
    assert all(len(r[1]) <= N_LINES for r in P["rows"])
# every box has the height of N_LINES lines, so that the two panels are equal
BOX = HEAD_H + N_LINES * LH + 0.05
H_PANEL = 2 * MARGIN + IMG_H + GAP + 3 * BOX + 2 * GAP
H = H_PANEL + (0.0 if ONLY_B else SUB_H)

fig = plt.figure(figsize=(W, H)); ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, W); ax.set_ylim(H, 0); ax.axis("off")
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(W), Inches(H)
SH = prs.slides.add_slide(prs.slide_layouts[6]).shapes
Z = [1]
def z():
    Z[0] += 1; return Z[0]


def frame(x, y, w, h, lc=RULE, lw=FRAME_PT, fc=None):
    ax.add_patch(Rectangle((x, y), w, h, fc=rgb(fc) if fc else "none", ec=rgb(lc) if lc else "none", lw=lw, zorder=z()))
    s = SH.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fc: s.fill.solid(); s.fill.fore_color.rgb = RGB(fc)
    else: s.fill.background()
    if lc: s.line.color.rgb = RGB(lc); s.line.width = Pt(lw)
    else: s.line.fill.background()
    s.shadow.inherit = False


def rule(x1, y1, x2, y2, color=RULE, lw=RULE_PT):
    ax.plot([x1, x2], [y1, y2], color=rgb(color), lw=lw, zorder=z(), solid_capstyle="butt")
    c = SH.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGB(color); c.line.width = Pt(lw)


def label(x, y, w, h, s, size, color=INK, bold=False, italic=False, align="left"):
    tx = {"left": x, "right": x + w, "center": x + w / 2}[align]
    ax.text(tx, y + h / 2, s, fontsize=size, color=rgb(color), family=FONT, ha=align, va="center",
            fontweight="bold" if bold else "normal", fontstyle="italic" if italic else "normal", zorder=z())
    tb = SH.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[align]
    r = p.add_run(); r.text = s; f = r.font; f.size = Pt(size); f.name = FONT; f.bold = bold; f.italic = italic
    f.color.rgb = RGB(color)


def words(x, y, w, lines):
    sp = tw(" ", BODY_PT)
    for i, ln in enumerate(lines):
        cx, cy = x, y + i * LH
        for word in ln:
            for q, (t, st) in enumerate(word):
                w_ = tw(t, BODY_PT, st is not None)
                if st is not None:                   # inside a word the tint stops at the neighbouring letters
                    l_ = 0.006 if q == 0 else 0.0; r_ = 0.006 if q == len(word) - 1 else 0.0
                    ax.add_patch(Rectangle((cx - l_, cy + 0.012), w_ + l_ + r_, LH - 0.02, fc=rgb(STYLE[st][1]), ec="none", zorder=z()))
                ax.text(cx, cy + LH / 2, t, fontsize=BODY_PT, family=FONT, va="center", ha="left", zorder=z(),
                        color=rgb(STYLE[st][0] if st else INK), fontweight="bold" if st else "normal")
                cx += w_
            cx += sp
    tb = SH.add_textbox(Inches(x), Inches(y), Inches(w), Inches(LH * len(lines))); tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.line_spacing = Pt(LH * 72)
        for j, word in enumerate(ln):
            for q, (t, st) in enumerate(word):
                r = p.add_run(); r.text = t + (" " if (j < len(ln) - 1 and q == len(word) - 1) else ""); f = r.font
                f.size = Pt(BODY_PT); f.name = FONT; f.bold = st is not None; f.color.rgb = RGB(STYLE[st][0] if st else INK)
                if st is not None:
                    rPr = r._r.get_or_add_rPr(); hl = rPr.makeelement(qn("a:highlight"), {})
                    hl.append(rPr.makeelement(qn("a:srgbClr"), {"val": STYLE[st][1]}))
                    rPr.insert_element_before(hl, "a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill", "a:latin", "a:ea", "a:cs",
                                              "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")


def panel(XO, P, tag):
    X0, Y0 = XO + MARGIN, MARGIN
    # top strip: image and instruction
    tmp = os.path.join(HERE, f"_case_2col_{tag}.jpg"); P["image"].save(tmp, quality=95)
    ax.imshow(np.asarray(P["image"]), extent=(X0, X0 + IMG_W, Y0 + IMG_H, Y0), zorder=z(), aspect="auto")
    SH.add_picture(tmp, Inches(X0), Inches(Y0), Inches(IMG_W), Inches(IMG_H))
    frame(X0, Y0, IMG_W, IMG_H)
    tx0, tw0 = X0 + IMG_W + 0.10, CW_ - IMG_W - 0.10
    label(tx0, Y0 + 0.015, tw0, 0.14, "Instruction", HEAD_PT, bold=True)
    label(tx0, Y0 + 0.150, tw0, 0.14, "“Please describe the image in detail.”", BODY_PT, italic=True)
    anyd = any(len(r) > 3 and r[3] for r in P["rows"])
    ly = Y0 + IMG_H - (2 * LH + 0.02 if anyd else LH + 0.005)
    if anyd:
        ly2 = ly + LH + 0.01
        ax.add_patch(Rectangle((tx0, ly2 + 0.012), 0.205, LH - 0.02, fc=rgb(BLUE_BG), ec="none", zorder=z()))
        label(tx0 + 0.012, ly2, 0.2, LH, "00", BODY_PT, color=BLUE, bold=True)
        label(tx0 + 0.255, ly2, tw0 - 0.255, LH, "digits of a number written in the image", BODY_PT, color=NOTE)
    ax.add_patch(Rectangle((tx0, ly + 0.012), 0.205, LH - 0.02, fc=rgb(RED_BG), ec="none", zorder=z()))
    _s = SH.add_shape(MSO_SHAPE.RECTANGLE, Inches(tx0), Inches(ly + 0.012), Inches(0.205), Inches(LH - 0.02))
    _s.fill.solid(); _s.fill.fore_color.rgb = RGB(RED_BG); _s.line.fill.background(); _s.shadow.inherit = False
    label(tx0 + 0.012, ly, 0.2, LH, "the", BODY_PT, color=RED, bold=True)
    label(tx0 + 0.255, ly, tw0 - 0.255, LH, "token repeating the one before it", BODY_PT, color=NOTE)
    # the three responses
    y = Y0 + IMG_H + GAP
    for r in P["rows"]:
        name, lines, n = r[:3]; digit = r[3] if len(r) > 3 else False
        frame(X0, y, CW_, HEAD_H, lc=None, fc=HEAD_BG)
        frame(X0, y, CW_, BOX); rule(X0, y + HEAD_H, X0 + CW_, y + HEAD_H)
        label(X0 + PAD, y, 2.2, HEAD_H, name, HEAD_PT, bold=True)
        if n >= 2: note, col = "longest repeated run: %d tokens" % n, RED
        elif digit: note, col = "repeats: the digits of the number only", BLUE
        else: note, col = "no repetition", GREEN
        label(X0 + CW_ - PAD - 1.9, y, 1.9, HEAD_H, note, BODY_PT, color=col, italic=True, bold=(n < 2), align="right")
        words(X0 + PAD, y + HEAD_H + 0.025, TEXT_W, lines)
        y += BOX + GAP
    label(XO, H_PANEL - 0.01, WP, SUB_H, P["sub"], SUB_PT, align="center")


if not ONLY_B: panel(0.0, PANEL_A, "a")
panel(0.0 if ONLY_B else WP + SEP, PANEL_B, "b_" + CASE)

if PREVIEW:
    base = os.path.join(HERE, ("preview_case_1col_" if ONLY_B else "preview_case_2col_") + CASE)
else:
    base = os.path.join(PAPER, "images", "fig-case-study-" + CASE if ONLY_B else "fig-case-study-2col")
fig.savefig(base + ".pdf"); fig.savefig(base + ".png", dpi=400)
prs.save(os.path.join(HERE, f"CoTA++_case_study_2col_{CASE}.pptx"))
print("written: %s.pdf/.png (%.3f x %.3f in) and _tools/CoTA++_case_study_2col_%s.pptx" % (base, W, H, CASE))
for P in (PANEL_A, PANEL_B):
    for r in P["rows"]:
        name, lines, n = r[:3]
        print("  %-24s run %2d | %d lines" % (name, n, len(lines)))
        for ln in lines: print("      " + " ".join("".join(t for t, _ in w_) for w_ in ln))
